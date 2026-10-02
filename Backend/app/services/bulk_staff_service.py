import io
import logging
import re
from datetime import date, datetime
from typing import Any
from uuid import UUID

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.exceptions import BadRequestException
from app.models.enums import Gender
from app.models.school import School
from app.models.smtp_configuration import SmtpConfiguration
from app.models.staff import StaffProfile
from app.models.user import User
from app.schemas.profile import StaffProfileCreate
from app.schemas.staff import (
    BulkStaffCreationResponse,
    BulkStaffCreationResultItem,
    BulkStaffPreviewResponse,
    BulkStaffPreviewRow,
    BulkStaffRowData,
    CreateStaffRequest,
)
from app.services.email_service import send_staff_invitation_email
from app.services.user_service import UserService

logger = logging.getLogger(__name__)

MAX_BULK_IMPORT_ROWS = 500
PHONE_REGEX = re.compile(r"^[0-9]{10,15}$")
EMAIL_REGEX = re.compile(r"^[\w\.-]+@[\w\.-]+\.\w+$")

# 6 Columns Contract: father names & roll_no removed; staff ID auto-generated sequentially
EXPECTED_COLUMNS = [
    "first_name",
    "last_name",
    "login_mobile",
    "email",
    "gender",
    "date_of_birth",
]

COLUMNS_METADATA = [
    {
        "name": "first_name",
        "label": "First Name",
        "description": "Staff member's first name (1-100 characters)",
        "example": "Rajesh",
        "required": True,
    },
    {
        "name": "last_name",
        "label": "Last Name",
        "description": "Staff member's last name (1-100 characters)",
        "example": "Sharma",
        "required": True,
    },
    {
        "name": "login_mobile",
        "label": "Login Mobile",
        "description": "10-15 digit mobile number without spaces or symbols",
        "example": "9876543210",
        "required": True,
    },
    {
        "name": "email",
        "label": "Email Address",
        "description": "Unique email address for account & invitations",
        "example": "rajesh.sharma@school.edu",
        "required": True,
    },
    {
        "name": "gender",
        "label": "Gender",
        "description": "Text value: 'Male', 'Female', or 'Other'",
        "example": "Male",
        "required": True,
    },
    {
        "name": "date_of_birth",
        "label": "Date of Birth",
        "description": "Date of birth in YYYY-MM-DD format (e.g. 1988-06-15)",
        "example": "1988-06-15",
        "required": True,
    },
]

MOCK_STAFF_RECORDS = [
    {
        "first_name": "Rajesh",
        "last_name": "Sharma",
        "login_mobile": "9876543210",
        "email": "rajesh.sharma@school.edu",
        "gender": "Male",
        "date_of_birth": "1988-06-15",
    },
    {
        "first_name": "Priya",
        "last_name": "Patel",
        "login_mobile": "9876543211",
        "email": "priya.patel@school.edu",
        "gender": "Female",
        "date_of_birth": "1992-09-22",
    },
    {
        "first_name": "Amit",
        "last_name": "Verma",
        "login_mobile": "9876543212",
        "email": "amit.verma@school.edu",
        "gender": "Male",
        "date_of_birth": "1985-12-05",
    },
]


def generate_sample_xlsx() -> bytes:
    """
    Generate an official sample XLSX template dynamically from the 6-column specification.
    Includes clean header styling, 3 realistic mock data rows, and an Instructions worksheet.
    """
    wb = openpyxl.Workbook()

    # ----------------------------------------------------
    # Sheet 1: Staff_Import_Template
    # ----------------------------------------------------
    ws_data = wb.active
    ws_data.title = "Staff_Import_Template"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Write Headers
    for col_idx, col_name in enumerate(EXPECTED_COLUMNS, start=1):
        cell = ws_data.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws_data.row_dimensions[1].height = 28

    # Write 3 Mock Rows
    for row_idx, record in enumerate(MOCK_STAFF_RECORDS, start=2):
        ws_data.row_dimensions[row_idx].height = 22
        for col_idx, col_name in enumerate(EXPECTED_COLUMNS, start=1):
            val = record.get(col_name)
            cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_name in ("gender", "login_mobile", "date_of_birth"):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths
    for col in ws_data.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell_val = str(cell.value or "")
            if len(cell_val) > max_len:
                max_len = len(cell_val)
        ws_data.column_dimensions[col_letter].width = max(max_len + 5, 18)

    # ----------------------------------------------------
    # Sheet 2: Instructions
    # ----------------------------------------------------
    ws_instr = wb.create_sheet(title="Instructions")
    ws_instr.row_dimensions[1].height = 28

    instr_headers = ["Column Name", "Required", "Description", "Accepted Format / Values", "Example"]
    for col_idx, h_text in enumerate(instr_headers, start=1):
        cell = ws_instr.cell(row=1, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    instructions_data = [
        ("first_name", "Yes", "Staff member's first name", "1-100 characters text", "Rajesh"),
        ("last_name", "Yes", "Staff member's last name", "1-100 characters text", "Sharma"),
        ("login_mobile", "Yes", "Mobile phone number used for login", "10-15 digits only, no spaces or special symbols", "9876543210"),
        ("email", "Yes", "Official email address (used for invitation link)", "Valid email format, must be unique across all users", "rajesh.sharma@school.edu"),
        ("gender", "Yes", "Staff member gender", "Text value: 'Male', 'Female', or 'Other'", "Male"),
        ("date_of_birth", "Yes", "Date of birth", "YYYY-MM-DD format", "1988-06-15"),
    ]

    for row_idx, item in enumerate(instructions_data, start=2):
        ws_instr.row_dimensions[row_idx].height = 24
        for col_idx, val in enumerate(item, start=1):
            cell = ws_instr.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 2, 5):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Add general notes below the instructions table
    start_note_row = len(instructions_data) + 4
    notes = [
        "IMPORTANT RULES & GUIDELINES:",
        "1. Do not rename, reorder, or remove any columns in the 'Staff_Import_Template' sheet.",
        "2. All 6 columns are mandatory. Blank or missing values in any row will cause that row to be rejected.",
        "3. Gender must be specified as 'Male', 'Female', or 'Other'.",
        "4. Staff IDs are automatically generated by the system in sequential order (e.g., 2026STF001). Do not include roll_no in this file.",
        "5. Passwords and PINs are NOT collected in this file.",
        "   Each created staff member will receive a secure invitation link to create their own password and 4-digit PIN.",
        "6. Duplicate emails or mobile numbers within this file or in the database will be rejected.",
        "7. Only valid rows will be imported. Any invalid or duplicate rows will be skipped safely.",
    ]

    for i, note in enumerate(notes):
        curr_row = start_note_row + i
        ws_instr.cell(row=curr_row, column=1, value=note)
        ws_instr.cell(row=curr_row, column=1).font = Font(
            name="Calibri", size=10, bold=(i == 0), color="0F172A" if i == 0 else "475569"
        )

    for col in ws_instr.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row <= len(instructions_data) + 2:
                cell_val = str(cell.value or "")
                if len(cell_val) > max_len:
                    max_len = len(cell_val)
        ws_instr.column_dimensions[col_letter].width = max(max_len + 4, 18)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalize_cell_value(val: Any) -> str:
    """Normalize cell values, removing trailing .0 from numeric floats/ints."""
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val)).strip()
        return str(val).strip()
    if isinstance(val, int):
        return str(val).strip()
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    return str(val).strip()


def parse_and_validate_file(
    file_bytes: bytes,
    school_id: UUID | None,
    admin_id: UUID,
    db: Session,
) -> BulkStaffPreviewResponse:
    """
    Parse an uploaded XLSX file, validate headers and every row, check in-file and database duplicates.
    Returns a complete dry-run / preview breakdown without making any changes to the database.
    """
    if not file_bytes:
        raise BadRequestException("The uploaded file is empty.")

    if len(file_bytes) > 10 * 1024 * 1024:
        raise BadRequestException("File size exceeds maximum allowed limit of 10 MB.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        logger.warning(f"Failed to read XLSX file: {exc}")
        raise BadRequestException("Invalid or corrupt XLSX file. Please ensure you upload a valid Excel (.xlsx) file.")

    if not wb.sheetnames:
        raise BadRequestException("The uploaded workbook contains no worksheets.")

    # Target the first worksheet
    ws = wb.active

    # Validate Header row (Row 1)
    header_row = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    while header_row and header_row[-1] is None:
        header_row.pop()

    if not header_row or all(h is None or str(h).strip() == "" for h in header_row):
        raise BadRequestException("Invalid file format. The uploaded file is missing the header row.")

    raw_headers = [str(h) if h is not None else "" for h in header_row]

    missing_cols = []
    unexpected_cols = []
    casing_errors = []

    expected_set = set(EXPECTED_COLUMNS)
    seen_headers = set()
    duplicate_headers = []

    for h in raw_headers:
        h_clean = h.strip()
        if not h_clean:
            continue
        if h_clean in seen_headers:
            duplicate_headers.append(h_clean)
        seen_headers.add(h_clean)

        if h != h_clean:
            casing_errors.append(f"Header '{h}' contains leading or trailing spaces.")
        elif h != h.lower():
            casing_errors.append(f"Header '{h}' must be all lowercase.")

        if h_clean not in expected_set:
            unexpected_cols.append(h_clean)

    for exp in EXPECTED_COLUMNS:
        if exp not in seen_headers:
            missing_cols.append(exp)

    if duplicate_headers:
        raise BadRequestException(f"Invalid file format. Duplicate column headers detected: {', '.join(duplicate_headers)}.")

    if missing_cols or unexpected_cols or casing_errors:
        err_msg_parts = ["Invalid file format. The uploaded file does not match the Staff Import template."]
        if missing_cols:
            err_msg_parts.append(f"Missing columns: {', '.join(missing_cols)}")
        if unexpected_cols:
            err_msg_parts.append(f"Unexpected columns: {', '.join(unexpected_cols)}")
        if casing_errors:
            err_msg_parts.append(f"Header formatting issues: {'; '.join(casing_errors)}")
        raise BadRequestException("\n".join(err_msg_parts))

    # Map column names to column indexes
    col_index_map = {str(h).strip(): idx for idx, h in enumerate(raw_headers, start=1)}

    # Read data rows (Row 2 onwards, numbering data records starting from 1)
    raw_data_rows: list[tuple[int, dict[str, Any]]] = []
    data_row_counter = 0
    for r in range(2, ws.max_row + 1):
        row_dict = {}
        has_content = False
        for col_name in EXPECTED_COLUMNS:
            col_idx = col_index_map[col_name]
            raw_val = ws.cell(row=r, column=col_idx).value
            normalized_val = _normalize_cell_value(raw_val)
            row_dict[col_name] = normalized_val
            if normalized_val:
                has_content = True

        if has_content:
            data_row_counter += 1
            raw_data_rows.append((data_row_counter, row_dict))

    if not raw_data_rows:
        raise BadRequestException("The uploaded file does not contain any staff records.")

    if len(raw_data_rows) > MAX_BULK_IMPORT_ROWS:
        raise BadRequestException(
            f"The file contains {len(raw_data_rows)} rows. Maximum allowed limit is {MAX_BULK_IMPORT_ROWS} staff records per import."
        )

    # Pre-fetch existing database identifiers for batch duplicate detection
    file_emails = {r["email"].lower() for _, r in raw_data_rows if r["email"]}
    file_mobiles = {r["login_mobile"] for _, r in raw_data_rows if r["login_mobile"]}

    # Existing emails in the system (users table)
    existing_db_emails: set[str] = set()
    if file_emails:
        stmt_email = select(func.lower(User.email)).where(
            User.deleted_at.is_(None),
            func.lower(User.email).in_(file_emails),
        )
        existing_db_emails = {e.lower() for e in db.scalars(stmt_email).all() if e}

    # Existing mobiles in the system (users table)
    existing_db_mobiles: set[str] = set()
    if file_mobiles:
        stmt_mobile = select(User.login_mobile).where(
            User.deleted_at.is_(None),
            User.login_mobile.in_(file_mobiles),
        )
        existing_db_mobiles = {m for m in db.scalars(stmt_mobile).all() if m}

    # Check SMTP configuration
    smtp_config = (
        db.query(SmtpConfiguration)
        .filter(
            SmtpConfiguration.admin_id == admin_id,
            SmtpConfiguration.deleted_at.is_(None),
            SmtpConfiguration.is_active.is_(True),
        )
        .first()
    )
    is_smtp_configured = bool(smtp_config and smtp_config.smtp_host)

    # Track in-file duplicates
    seen_file_emails: dict[str, int] = {}
    seen_file_mobiles: dict[str, int] = {}

    preview_rows: list[BulkStaffPreviewRow] = []
    valid_count = 0
    duplicate_count = 0
    invalid_count = 0

    for row_num, row_data in raw_data_rows:
        errors: list[str] = []
        is_duplicate = False

        fn = row_data["first_name"]
        ln = row_data["last_name"]
        mobile = row_data["login_mobile"]
        email = row_data["email"]
        gender_raw = row_data["gender"]
        dob_raw = row_data["date_of_birth"]

        # 1. first_name validation
        if not fn:
            errors.append("First name is required")
        elif len(fn) > 100:
            errors.append("First name cannot exceed 100 characters")

        # 2. last_name validation
        if not ln:
            errors.append("Last name is required")
        elif len(ln) > 100:
            errors.append("Last name cannot exceed 100 characters")

        # 3. login_mobile validation
        if not mobile:
            errors.append("Mobile number is required")
        elif not PHONE_REGEX.match(mobile):
            errors.append("Mobile number must be 10 to 15 digits without spaces or symbols")
        else:
            if mobile in existing_db_mobiles:
                errors.append("Mobile number already exists in system")
                is_duplicate = True
            elif mobile in seen_file_mobiles:
                errors.append(f"Duplicate mobile number (same as row {seen_file_mobiles[mobile]})")
                is_duplicate = True
            else:
                seen_file_mobiles[mobile] = row_num

        # 4. email validation
        if not email:
            errors.append("Email address is required")
        elif len(email) > 100:
            errors.append("Email cannot exceed 100 characters")
        elif not EMAIL_REGEX.match(email):
            errors.append("Invalid email address format")
        else:
            email_lower = email.lower()
            if email_lower in existing_db_emails:
                errors.append("Email already exists in system")
                is_duplicate = True
            elif email_lower in seen_file_emails:
                errors.append(f"Duplicate email in file (same as row {seen_file_emails[email_lower]})")
                is_duplicate = True
            else:
                seen_file_emails[email_lower] = row_num

        # 5. gender validation (supports 'Male', 'Female', 'Other')
        gender_display = ""
        if not gender_raw:
            errors.append("Gender is required ('Male', 'Female', or 'Other')")
        else:
            g_str = str(gender_raw).strip().lower()
            if g_str in ("male", "m", "1"):
                gender_display = "Male"
            elif g_str in ("female", "f", "2"):
                gender_display = "Female"
            elif g_str in ("other", "o", "3"):
                gender_display = "Other"
            else:
                errors.append("Invalid gender. Must be 'Male', 'Female', or 'Other'")

        # 6. date_of_birth validation
        dob_normalized = ""
        if not dob_raw:
            errors.append("Date of birth is required")
        else:
            try:
                if isinstance(dob_raw, str):
                    parsed_date = datetime.strptime(dob_raw, "%Y-%m-%d").date()
                elif isinstance(dob_raw, (datetime, date)):
                    parsed_date = dob_raw if isinstance(dob_raw, date) else dob_raw.date()
                else:
                    parsed_date = datetime.strptime(str(dob_raw), "%Y-%m-%d").date()

                today = datetime.now().date()
                if parsed_date >= today:
                    errors.append("Date of birth must be a past date")
                elif parsed_date.year < 1920:
                    errors.append("Date of birth must be after year 1920")
                else:
                    dob_normalized = parsed_date.strftime("%Y-%m-%d")
            except Exception:
                errors.append("Invalid date of birth. Must be a valid date in YYYY-MM-DD format")

        # Determine row status
        if is_duplicate:
            status = "DUPLICATE"
            duplicate_count += 1
        elif errors:
            status = "INVALID"
            invalid_count += 1
        else:
            status = "VALID"
            valid_count += 1

        row_obj = BulkStaffPreviewRow(
            row_number=row_num,
            data=BulkStaffRowData(
                first_name=fn,
                last_name=ln,
                login_mobile=mobile,
                email=email,
                gender=gender_display or str(gender_raw),
                date_of_birth=dob_normalized or str(dob_raw),
            ),
            status=status,
            errors=errors,
        )
        preview_rows.append(row_obj)

    return BulkStaffPreviewResponse(
        total_rows=len(preview_rows),
        valid_rows=valid_count,
        duplicate_rows=duplicate_count,
        invalid_rows=invalid_count,
        is_smtp_configured=is_smtp_configured,
        rows=preview_rows,
    )


def execute_bulk_import(
    file_bytes: bytes,
    current_user: User,
    db: Session,
    user_service: UserService,
) -> BulkStaffCreationResponse:
    """
    Authoritatively execute bulk staff creation for valid rows.
    Staff IDs are automatically generated in sequential order internally.
    Re-validates the file, re-checks database uniqueness to prevent race conditions,
    creates accounts using UserService.create_staff atomically per record,
    generates setup tokens, sends invitations if SMTP configured, and aggregates results.
    """
    school_id = current_user.school_id

    # 1. Authoritative validation pass
    preview = parse_and_validate_file(
        file_bytes=file_bytes,
        school_id=school_id,
        admin_id=current_user.id,
        db=db,
    )

    # 2. Lookup school name and SMTP configuration
    school_name = "School LMS"
    if school_id:
        school = db.query(School).filter(School.id == school_id, School.deleted_at.is_(None)).first()
        if school and school.name:
            school_name = school.name

    smtp_config = (
        db.query(SmtpConfiguration)
        .filter(
            SmtpConfiguration.admin_id == current_user.id,
            SmtpConfiguration.deleted_at.is_(None),
            SmtpConfiguration.is_active.is_(True),
        )
        .first()
    )
    is_smtp_configured = bool(smtp_config and smtp_config.smtp_host)

    results: list[BulkStaffCreationResultItem] = []
    created_count = 0
    failed_count = 0
    duplicate_count = preview.duplicate_rows
    invalid_count = preview.invalid_rows
    email_sent_count = 0
    email_failed_count = 0

    for row in preview.rows:
        # Non-valid rows are passed directly to results with their errors
        if row.status != "VALID":
            results.append(
                BulkStaffCreationResultItem(
                    row_number=row.row_number,
                    staff_id=None,
                    roll_no=None,
                    name=f"{row.data.first_name} {row.data.last_name}".strip(),
                    email=row.data.email,
                    login_mobile=row.data.login_mobile,
                    status=row.status,
                    created=False,
                    email_sent=False,
                    errors=row.errors,
                )
            )
            continue

        # Re-check database duplicates immediately before creation for race safety
        try:
            # Check email
            existing_email = (
                db.query(User.id)
                .filter(
                    User.deleted_at.is_(None),
                    func.lower(User.email) == row.data.email.lower(),
                )
                .first()
            )
            if existing_email:
                duplicate_count += 1
                results.append(
                    BulkStaffCreationResultItem(
                        row_number=row.row_number,
                        staff_id=None,
                        roll_no=None,
                        name=f"{row.data.first_name} {row.data.last_name}".strip(),
                        email=row.data.email,
                        login_mobile=row.data.login_mobile,
                        status="DUPLICATE",
                        created=False,
                        email_sent=False,
                        errors=["Email already exists in system"],
                    )
                )
                continue

            # Check mobile
            existing_mobile = (
                db.query(User.id)
                .filter(
                    User.deleted_at.is_(None),
                    User.login_mobile == row.data.login_mobile,
                )
                .first()
            )
            if existing_mobile:
                duplicate_count += 1
                results.append(
                    BulkStaffCreationResultItem(
                        row_number=row.row_number,
                        staff_id=None,
                        roll_no=None,
                        name=f"{row.data.first_name} {row.data.last_name}".strip(),
                        email=row.data.email,
                        login_mobile=row.data.login_mobile,
                        status="DUPLICATE",
                        created=False,
                        email_sent=False,
                        errors=["Mobile number already exists in system"],
                    )
                )
                continue

            # Parse date of birth & gender
            parsed_dob = datetime.strptime(row.data.date_of_birth, "%Y-%m-%d").date()
            g_lower = row.data.gender.lower()
            gender_enum = (
                Gender.MALE if g_lower in ("male", "m", "1")
                else Gender.FEMALE if g_lower in ("female", "f", "2")
                else Gender.OTHER
            )

            # Construct typed CreateStaffRequest with roll_no="" so UserService assigns sequential Staff ID
            create_req = CreateStaffRequest(
                first_name=row.data.first_name,
                last_name=row.data.last_name,
                login_mobile=row.data.login_mobile,
                email=row.data.email,
                profile=StaffProfileCreate(
                    roll_no="",
                    gender=gender_enum,
                    date_of_birth=parsed_dob,
                ),
            )

            # Atomically create user, profile (with sequential staff ID), and onboarding token
            user, raw_token = user_service.create_staff(
                data=create_req,
                created_by_id=current_user.id,
                school_id=school_id,
            )

            assigned_roll_no = user.staff_profile.roll_no if user.staff_profile else None
            setup_url = f"{settings.FRONTEND_URL}/login/setup?token={raw_token}"
            email_sent = False
            email_error = None

            # Attempt email invitation if SMTP is configured
            if is_smtp_configured:
                try:
                    staff_name = f"{user.first_name} {user.last_name}".strip()
                    email_sent = send_staff_invitation_email(
                        staff_name=staff_name,
                        staff_email=user.email,
                        school_name=school_name,
                        raw_token=raw_token,
                        smtp_config=smtp_config,
                    )
                    if email_sent:
                        email_sent_count += 1
                    else:
                        email_failed_count += 1
                except Exception as exc:
                    logger.error(f"Bulk import email send failed for {user.email}: {exc}")
                    email_sent = False
                    email_error = str(exc)
                    email_failed_count += 1
            else:
                email_failed_count += 1

            created_count += 1
            results.append(
                BulkStaffCreationResultItem(
                    row_number=row.row_number,
                    staff_id=user.id,
                    roll_no=assigned_roll_no,
                    name=f"{user.first_name} {user.last_name}".strip(),
                    email=user.email,
                    login_mobile=user.login_mobile,
                    status="CREATED",
                    created=True,
                    email_sent=email_sent,
                    email_error=email_error,
                    setup_url=setup_url,
                    errors=[],
                )
            )

        except Exception as exc:
            db.rollback()
            logger.error(f"Unexpected failure creating staff row {row.row_number}: {exc}", exc_info=True)
            failed_count += 1
            results.append(
                BulkStaffCreationResultItem(
                    row_number=row.row_number,
                    staff_id=None,
                    roll_no=None,
                    name=f"{row.data.first_name} {row.data.last_name}".strip(),
                    email=row.data.email,
                    login_mobile=row.data.login_mobile,
                    status="FAILED",
                    created=False,
                    email_sent=False,
                    errors=[f"Creation failed: {str(exc)}"],
                )
            )

    logger.info(
        f"Bulk staff import finished by user {current_user.id} in school {school_id}: "
        f"{created_count} created, {duplicate_count} duplicates, {invalid_count} invalid, {failed_count} failed."
    )

    msg = f"Bulk import complete. {created_count} staff accounts created successfully."
    if failed_count > 0:
        msg += f" {failed_count} records failed."

    return BulkStaffCreationResponse(
        total_rows=len(preview.rows),
        created_count=created_count,
        failed_count=failed_count,
        duplicate_count=duplicate_count,
        invalid_count=invalid_count,
        email_sent_count=email_sent_count,
        email_failed_count=email_failed_count,
        results=results,
        message=msg,
    )
