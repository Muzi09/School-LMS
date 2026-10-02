import io
import openpyxl
from datetime import date, datetime, timedelta, timezone
from uuid import uuid4

from app.core.database import SessionLocal
from app.models.enums import Gender, UserRole
from app.models.principal import PrincipalProfile
from app.models.school import School
from app.models.staff import StaffProfile
from app.models.user import User
from app.repositories.user_repository import UserRepository
from app.schemas.profile import StaffProfileCreate
from app.schemas.staff import CreateStaffRequest
from app.services.bulk_staff_service import (
    EXPECTED_COLUMNS,
    execute_bulk_import,
    generate_sample_xlsx,
    parse_and_validate_file,
)
from app.services.user_service import UserService


def create_test_xlsx(rows: list[dict], headers: list[str] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    hdrs = headers if headers is not None else EXPECTED_COLUMNS
    for col_idx, h in enumerate(hdrs, start=1):
        ws.cell(row=1, column=col_idx, value=h)

    for r_idx, row_dict in enumerate(rows, start=2):
        for col_idx, h in enumerate(hdrs, start=1):
            ws.cell(row=r_idx, column=col_idx, value=row_dict.get(h, ""))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_tests():
    db = SessionLocal()
    user_repo = UserRepository(db)
    user_service = UserService(user_repo=user_repo, db=db)

    print("Running Bulk Staff Import Backend Tests...")

    # 1. Test sample xlsx generation (6 columns only)
    sample_bytes = generate_sample_xlsx()
    assert len(sample_bytes) > 0, "Sample bytes should not be empty"
    wb_sample = openpyxl.load_workbook(io.BytesIO(sample_bytes))
    assert "Staff_Import_Template" in wb_sample.sheetnames, "Template sheet missing"
    assert "Instructions" in wb_sample.sheetnames, "Instructions sheet missing"
    ws = wb_sample["Staff_Import_Template"]
    read_headers = [ws.cell(row=1, column=c).value for c in range(1, len(EXPECTED_COLUMNS) + 1)]
    assert read_headers == EXPECTED_COLUMNS, f"Headers mismatch: {read_headers} != {EXPECTED_COLUMNS}"
    assert len(EXPECTED_COLUMNS) == 6, f"Expected 6 columns, got {len(EXPECTED_COLUMNS)}"
    assert ws.max_row >= 4, "Should have header + 3 mock rows"
    print("[PASS] Test 1: Sample XLSX Generation (6 columns) Passed")

    # 2. Setup mock School and Admin User
    school = School(
        name=f"Bulk Test School {uuid4().hex[:6]}",
        code=f"BTS{uuid4().hex[:4].upper()}",
    )
    db.add(school)
    db.commit()
    db.refresh(school)

    admin_user = User(
        first_name="Admin",
        last_name="Tester",
        email=f"admin_{uuid4().hex[:6]}@school.edu",
        login_mobile=f"99{int(uuid4().int % 100000000):08d}",
        role=UserRole.PRINCIPAL,
        is_active=True,
    )
    db.add(admin_user)
    db.commit()
    db.refresh(admin_user)

    admin_profile = PrincipalProfile(
        user_id=admin_user.id,
        school_id=school.id,
    )
    db.add(admin_profile)
    db.commit()
    db.refresh(admin_user)

    # 3. Test Invalid Headers (Missing and Extra columns)
    bad_headers_xlsx = create_test_xlsx(
        rows=[{"first_name": "Test"}],
        headers=["first_name", "last_name", "login_mobile", "email", "extra_col"],
    )
    try:
        parse_and_validate_file(bad_headers_xlsx, school_id=school.id, admin_id=admin_user.id, db=db)
        assert False, "Should have failed with invalid headers"
    except Exception as exc:
        assert "Invalid file format" in str(exc)
        assert "Missing columns" in str(exc)
        assert "Unexpected columns" in str(exc)
    print("[PASS] Test 2: Header Validation (Missing & Unexpected Columns) Passed")

    # 4. Test In-File Duplicates
    dup_mobile = f"98{int(uuid4().int % 100000000):08d}"
    dup_email = f"duplicate_{uuid4().hex[:4]}@school.edu"

    test_rows_with_dups = [
        {
            "first_name": "Alice",
            "last_name": "Smith",
            "login_mobile": dup_mobile,
            "email": dup_email,
            "gender": "Female",
            "date_of_birth": "1990-01-15",
        },
        {
            "first_name": "Bob",
            "last_name": "Smith",
            "login_mobile": dup_mobile,  # Duplicate mobile
            "email": dup_email,  # Duplicate email
            "gender": "Male",
            "date_of_birth": "1988-05-20",
        },
    ]
    dup_xlsx = create_test_xlsx(test_rows_with_dups)
    preview = parse_and_validate_file(dup_xlsx, school_id=school.id, admin_id=admin_user.id, db=db)
    assert preview.total_rows == 2
    assert preview.valid_rows == 1
    assert preview.duplicate_rows == 1
    assert preview.rows[0].row_number == 1, f"Expected row_number 1, got {preview.rows[0].row_number}"
    assert preview.rows[1].row_number == 2, f"Expected row_number 2, got {preview.rows[1].row_number}"
    assert preview.rows[1].status == "DUPLICATE"
    assert any("same as row 1" in err for err in preview.rows[1].errors), f"Expected error referencing row 1: {preview.rows[1].errors}"
    print("[PASS] Test 3: In-File Duplicate Detection & Row Numbering starting at 1 Passed")

    # 5. Test Database Duplicates
    existing_staff_email = f"existing_{uuid4().hex[:6]}@school.edu"
    existing_staff_mobile = f"91{int(uuid4().int % 100000000):08d}"

    existing_user, _ = user_service.create_staff(
        data=CreateStaffRequest(
            first_name="Existing",
            last_name="Staff",
            email=existing_staff_email,
            login_mobile=existing_staff_mobile,
            profile=StaffProfileCreate(
                roll_no="",
                gender=Gender.FEMALE,
                date_of_birth=date(1992, 3, 10),
            ),
        ),
        created_by_id=admin_user.id,
        school_id=school.id,
    )
    assert existing_user is not None

    test_db_dup_rows = [
        {
            "first_name": "Clara",
            "last_name": "Doe",
            "login_mobile": existing_staff_mobile,  # Exists in DB
            "email": existing_staff_email,  # Exists in DB
            "gender": "Female",
            "date_of_birth": "1994-07-21",
        },
    ]
    db_dup_xlsx = create_test_xlsx(test_db_dup_rows)
    db_dup_preview = parse_and_validate_file(db_dup_xlsx, school_id=school.id, admin_id=admin_user.id, db=db)
    assert db_dup_preview.duplicate_rows == 1
    assert db_dup_preview.valid_rows == 0
    assert any("already exists in system" in err for err in db_dup_preview.rows[0].errors)
    print("[PASS] Test 4: Database Duplicate Detection Passed")

    # 6. Test Row Validations (Invalid email, mobile, gender, date)
    invalid_rows = [
        {
            "first_name": "",  # Missing
            "last_name": "Invalid",
            "login_mobile": "123",  # Too short
            "email": "not-an-email",  # Bad format
            "gender": "unknown_gender",  # Bad gender text
            "date_of_birth": "2099-01-01",  # Future date
        }
    ]
    inv_xlsx = create_test_xlsx(invalid_rows)
    inv_preview = parse_and_validate_file(inv_xlsx, school_id=school.id, admin_id=admin_user.id, db=db)
    assert inv_preview.invalid_rows == 1
    assert inv_preview.valid_rows == 0
    row_errs = inv_preview.rows[0].errors
    assert any("First name is required" in e for e in row_errs)
    assert any("Mobile number must be" in e for e in row_errs)
    assert any("Invalid email" in e for e in row_errs)
    assert any("Invalid gender" in e for e in row_errs)
    assert any("past date" in e for e in row_errs)
    print("[PASS] Test 5: Row Field Validation Constraints (Textual gender) Passed")

    # 7. Test Bulk Creation with Valid Rows & Sequential Staff ID Assignment
    valid_unique_1 = f"val1_{uuid4().hex[:4]}"
    valid_unique_2 = f"val2_{uuid4().hex[:4]}"

    valid_creation_rows = [
        {
            "first_name": "David",
            "last_name": "Miller",
            "login_mobile": f"98{int(uuid4().int % 100000000):08d}",
            "email": f"{valid_unique_1}@school.edu",
            "gender": "Male",
            "date_of_birth": "1989-11-12",
        },
        {
            "first_name": "Emma",
            "last_name": "Watson",
            "login_mobile": f"97{int(uuid4().int % 100000000):08d}",
            "email": f"{valid_unique_2}@school.edu",
            "gender": "Female",
            "date_of_birth": "1991-04-15",
        },
    ]
    valid_xlsx = create_test_xlsx(valid_creation_rows)
    creation_res = execute_bulk_import(
        file_bytes=valid_xlsx,
        current_user=admin_user,
        db=db,
        user_service=user_service,
    )
    assert creation_res.total_rows == 2
    assert creation_res.created_count == 2
    assert creation_res.failed_count == 0
    assert len(creation_res.results) == 2

    # Verify sequential auto-generation of Staff IDs
    res1 = creation_res.results[0]
    res2 = creation_res.results[1]
    assert res1.created is True
    assert res2.created is True
    assert res1.roll_no is not None
    assert res2.roll_no is not None
    assert res1.roll_no != res2.roll_no

    current_year = datetime.now(timezone.utc).year
    prefix = f"{current_year}STF"
    assert res1.roll_no.startswith(prefix)
    assert res2.roll_no.startswith(prefix)

    # Extract sequences to confirm sequential generation
    seq1 = int(res1.roll_no[len(prefix):])
    seq2 = int(res2.roll_no[len(prefix):])
    assert seq2 == seq1 + 1, f"Expected consecutive IDs, got {res1.roll_no} and {res2.roll_no}"

    # Verify in database
    db_user_1 = user_repo.get_by_id_with_profiles(res1.staff_id)
    assert db_user_1.staff_profile.roll_no == res1.roll_no
    assert db_user_1.staff_profile.status == "PENDING_ACTIVATION"
    assert db_user_1.password_hash is None
    assert db_user_1.staff_profile.pin_hash is None

    print(f"[PASS] Test 6: Bulk Creation with Sequential Staff IDs ({res1.roll_no}, {res2.roll_no}) Passed")

    print("\nALL BULK STAFF IMPORT TESTS PASSED SUCCESSFULLY!")
    db.close()


if __name__ == "__main__":
    run_tests()
